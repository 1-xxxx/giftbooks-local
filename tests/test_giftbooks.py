"""Regression tests using real Flask/Excel/Pillow and a stub for Mac-only MLX.

Run: python -m unittest discover -s tests -p 'test_*.py' -v
"""
import base64
import importlib.util
import io
import sys
import tempfile
import threading
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook
import requests


def load_app():
    stubs = {name: types.ModuleType(name) for name in
             ('mlx_vlm', 'mlx_vlm.prompt_utils', 'mlx_vlm.utils')}
    stubs['mlx_vlm'].load = lambda name: (object(), object())
    stubs['mlx_vlm'].generate = lambda *args, **kwargs: '{}'
    stubs['mlx_vlm.prompt_utils'].apply_chat_template = lambda *args, **kwargs: 'prompt'
    stubs['mlx_vlm.utils'].load_config = lambda *args: {}
    with patch.dict(sys.modules, stubs):
        spec = importlib.util.spec_from_file_location('giftbooks_test_app', Path(__file__).parents[1] / 'giftbooks_local_v2.py')
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module


app = load_app()


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.root_patch = patch.object(app, 'DATA_ROOT', self.root)
        self.path_patch = patch.object(app, 'RESULTS_WORKBOOK_PATH', self.root / 'results.xlsx')
        self.root_patch.start()
        self.path_patch.start()
        self.client = app.app.test_client()
        self.metadata = app.SearchMetadata(title='Example book', authors=['Jane Author'])
        self.result = {'status': 'found', 'record': {'record_url': app.illinois_record_url('test')},
                       'search_url': app.illinois_search_url(self.metadata)}

    def tearDown(self):
        self.root_patch.stop()
        self.path_patch.stop()
        self.temp.cleanup()

    def save(self, **kwargs):
        return app.save_scan_result(self.metadata, self.result, 'keep', '', 12.3, **kwargs)

    def test_new_catalog_urls_save_and_finalize_same_row(self):
        saved = self.save()
        response = self.client.post('/api/finalize-time', json={'scan_id': saved['scan_id'], 'elapsed_seconds': 48.27})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(app.RESULTS_WORKBOOK_PATH)
        sheet = workbook[app.RESULTS_SHEET]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['O2'].value, 48.3)
        self.assertEqual(sheet['L2'].hyperlink.target, self.result['record']['record_url'])
        workbook.close()

    def test_legacy_catalog_url_remains_valid(self):
        self.result['record']['record_url'] = app.ILLINOIS_CATALOG_BASE + '/discovery/fulldisplay?docid=test'
        self.save()

    def test_different_edition_saves_search_url(self):
        self.result['status'] = 'different_edition'
        self.save()
        workbook = load_workbook(app.RESULTS_WORKBOOK_PATH)
        self.assertEqual(workbook[app.RESULTS_SHEET]['L2'].value, self.result['search_url'])
        workbook.close()

    def test_untrusted_catalog_host_rejected(self):
        self.result['record']['record_url'] = 'https://example.com/nde/fulldisplay?docid=test'
        with self.assertRaises(ValueError):
            self.save()
        self.assertFalse(app.RESULTS_WORKBOOK_PATH.exists())

    def test_transcript_and_audio_saved_next_to_scan(self):
        conversation = {'transcript': [{'text': 'Let us keep this book', 'at_seconds': 22.1}],
                        'audio': 'data:audio/webm;codecs=opus;base64,' + base64.b64encode(b'audio fixture').decode()}
        saved = self.save(conversation=conversation)
        self.assertEqual((self.root / 'conversations' / (saved['scan_id'] + '.webm')).read_bytes(), b'audio fixture')
        self.assertIn('Let us keep this book', (self.root / 'conversations' / (saved['scan_id'] + '.json')).read_text())

    def test_invalid_recording_does_not_create_row(self):
        for conversation in ({'audio': 'data:text/html;base64,eA=='}, {'transcript': ['bad']}, ['bad']):
            with self.subTest(conversation=conversation), self.assertRaises(ValueError):
                self.save(conversation=conversation)
        self.assertFalse(app.RESULTS_WORKBOOK_PATH.exists())

    def test_failed_excel_save_removes_recording_and_allows_retry(self):
        conversation = {'transcript': [{'text': 'keep'}]}
        with patch.object(app.os, 'replace', side_effect=PermissionError('Workbook open')):
            with self.assertRaises(PermissionError):
                self.save(conversation=conversation)
        self.assertEqual(list((self.root / 'conversations').iterdir()), [])
        self.assertFalse(app.RESULTS_WORKBOOK_PATH.exists())
        self.save(conversation=conversation)

    def test_invalid_and_unknown_timing_leave_workbook_unchanged(self):
        saved = self.save()
        for seconds in (-1, float('nan'), 86401, None):
            with self.subTest(seconds=seconds), self.assertRaises(ValueError):
                app.finalize_processing_time(saved['scan_id'], seconds)
        with self.assertRaises(ValueError):
            app.finalize_processing_time('scan_20260930_010000_000000', 50)
        workbook = load_workbook(app.RESULTS_WORKBOOK_PATH)
        self.assertEqual(workbook[app.RESULTS_SHEET]['O2'].value, 12.3)
        workbook.close()

    def test_finalization_failure_keeps_saved_row_for_retry(self):
        saved = self.save()
        with patch.object(app.os, 'replace', side_effect=PermissionError('Workbook open')):
            with self.assertRaises(PermissionError):
                app.finalize_processing_time(saved['scan_id'], 60)
        workbook = load_workbook(app.RESULTS_WORKBOOK_PATH)
        self.assertEqual(workbook[app.RESULTS_SHEET]['O2'].value, 12.3)
        workbook.close()
        app.finalize_processing_time(saved['scan_id'], 65)

    def test_searches_overlap_and_propagate_catalog_failure(self):
        metadata = app.SearchMetadata(title='Example', isbn13='9780306406157')
        barrier = threading.Barrier(2)

        def fetch(query):
            barrier.wait(timeout=2)
            return []

        with patch.object(app, 'fetch_illinois_catalog', side_effect=fetch) as mocked:
            self.assertEqual(app.search_illinois_catalog(metadata)['status'], 'not_found')
            self.assertEqual(mocked.call_count, 2)
        with patch.object(app, 'fetch_illinois_catalog', side_effect=requests.ConnectionError('Offline')):
            with self.assertRaises(requests.ConnectionError):
                app.search_illinois_catalog(metadata)

    def test_fast_image_size_and_traditional_size(self):
        image = Image.new('RGB', (2400, 1600), 'white')
        buffer = io.BytesIO()
        image.save(buffer, format='JPEG')
        photo = {'data': 'data:image/jpeg;base64,' + base64.b64encode(buffer.getvalue()).decode(), 'role': 'front_cover'}
        dimensions = []

        def generate(model, processor, prompt, paths, **kwargs):
            with Image.open(paths[0]) as normalized:
                dimensions.append(normalized.size)
            return '{"title":"Example"}'

        with patch.object(app, 'get_model_bundle', return_value=(None, None, {})), patch.object(app, 'generate', side_effect=generate):
            app.extract_search_metadata([photo], fast=True)
            app.extract_search_metadata([photo], fast=False)
        self.assertEqual(dimensions[0][0], app.FAST_IMAGE_SIZE)
        self.assertEqual(dimensions[1][0], 1800)

    def test_check_endpoint_forwards_mode_and_html_renders(self):
        with patch.object(app, 'extract_search_metadata', return_value=(self.metadata, '{}')) as extract, patch.object(app, 'search_illinois_catalog', return_value=self.result):
            response = self.client.post('/api/check', json={'images': [], 'mode': 'fast'})
            self.assertEqual(response.status_code, 200)
            extract.assert_called_once_with([], fast=True)
        html = self.client.get('/').data.decode()
        self.assertIn('Capture cover and check', html)
        self.assertNotIn('countdown', html)
        self.assertNotIn('{{', html)


if __name__ == '__main__':
    unittest.main()
